"""Data for the Grid view (one live card per share) and the Market Flow (tree) view, plus Excel/CSV import."""
from __future__ import annotations

import csv
import io
import re

from . import analytics as A
from .market import LOCK, MARKET


def _r(x, nd=2):
    return None if x is None else round(x, nd)


def _levels(c5, price, pdh, pdl, day_open):
    """Supports below / resistances above price from today's swing points (S1 nearest)."""
    sw = A.swings(c5, 2, 2) if len(c5) >= 5 else []
    lows = [s["price"] for s in sw if s["kind"] == "low" and s["price"] < price]
    highs = [s["price"] for s in sw if s["kind"] == "high" and s["price"] > price]
    for v in (pdl, day_open):
        if v and v < price:
            lows.append(v)
    for v in (pdh, day_open):
        if v and v > price:
            highs.append(v)

    def uniq(vals, rev):
        out = []
        for v in sorted(vals, reverse=rev):
            if not out or abs(v - out[-1]) / price > 0.0015:
                out.append(v)
        return out[:4]
    return uniq(lows, True), uniq(highs, False)


def card(engine, st) -> dict:
    now = MARKET.clock()
    with LOCK:
        today = [dict(c) for c in st.today_candles(now)]
        price = st.ltp
        pdh = st.daily[-1]["h"] if st.daily else None
        pdl = st.daily[-1]["l"] if st.daily else None
        day_open = st.open
    c5 = A.resample(today, 5)
    sig = engine.signals_for(st, 10)["intraday"] if price > 0 else {}
    sup, res = _levels(c5, price, pdh, pdl, day_open) if price > 0 else ([], [])
    lo = min([c["l"] for c in c5] + sup) if c5 else 0
    hi = max([c["h"] for c in c5] + res) if c5 else 0
    obs = [o for o in (sig.get("levels") or {}).get("order_blocks", []) if c5 and o["bottom"] <= hi and o["top"] >= lo]
    vw, pv, vv = [], 0.0, 0.0
    for c in c5:
        pv += (c["h"] + c["l"] + c["c"]) / 3 * c["v"]
        vv += c["v"]
        vw.append(_r(pv / vv) if vv else None)
    return {"sym": st.sym, "sector": st.sector,
            "candles": [[c["t"], c["o"], c["h"], c["l"], c["c"]] for c in c5],
            "vwap": vw, "obs": obs[:3], "S": [_r(x) for x in sup], "R": [_r(x) for x in res],
            "sig": {k: sig.get(k) for k in ("action", "entry", "sl", "t1", "t2", "rr1", "zone", "confidence", "summary")},
            "vwap_now": vw[-1] if vw else None}


def flow(syms: list[str]) -> dict:
    """Market Flow: every share's intraday path (% from previous close), sectors, bullish / bearish split."""
    now = MARKET.clock()
    open_t = A.ist_day_start(now) + A.SESSION_OPEN
    items, sectors = [], {}
    for s in syms:
        st = MARKET.stocks.get(s)
        if not st or not st.metrics:
            continue
        with LOCK:
            today = [dict(c) for c in st.today_candles(now)]
            m = dict(st.metrics)
        pc = m.get("prev_close") or st.prev_close or 0
        if not pc:
            continue
        c5 = A.resample(today, 5)
        path = [[int((c["t"] - open_t) // 300), round((c["c"] - pc) / pc * 100, 3)] for c in c5]
        last = c5[-24:]
        mv = max([c["v"] for c in last] or [1]) or 1
        vbars = [round(c["v"] / mv * (1 if c["c"] >= c["o"] else -1), 3) for c in last]
        items.append({"sym": s, "sector": st.sector, "chg": m["chg"], "ltp": m["ltp"], "strength": m["strength"],
                      "direction": m["direction"], "choppy": m["choppy"], "vol_ratio": m["vol_ratio"],
                      "pressure": m["pressure"], "path": path, "vbars": vbars})
        g = sectors.setdefault(st.sector, {"sector": st.sector, "n": 0, "sum": 0.0, "adv": 0, "dec": 0})
        g["n"] += 1
        g["sum"] += m["chg"]
        g["adv"] += m["chg"] > 0
        g["dec"] += m["chg"] < 0
    secs = sorted(({**g, "chg": round(g["sum"] / g["n"], 2)} for g in sectors.values() if g["n"]), key=lambda g: -g["chg"])
    bulls = sorted([i for i in items if i["direction"] > 0 and not i["choppy"]], key=lambda i: -i["strength"])
    bears = sorted([i for i in items if i["direction"] < 0 and not i["choppy"]], key=lambda i: -i["strength"])
    return {"items": items, "sectors": secs, "bull_count": sum(1 for i in items if i["chg"] > 0),
            "bear_count": sum(1 for i in items if i["chg"] < 0), "bulls": [i["sym"] for i in bulls[:16]],
            "bears": [i["sym"] for i in bears[:16]], "leading": secs[0] if secs else None,
            "lagging": secs[-1] if secs else None, "slots": (A.SESSION_CLOSE - A.SESSION_OPEN) // 300,
            "now_slot": max(0, (now - open_t) // 300), "ts": now}


# ------------------------------------------------------------------ Excel / CSV import
HEADER_RX = re.compile(r"symbol|stock|share|scrip|ticker|instrument|name|company|code", re.I)


def _clean(v: str) -> str:
    v = str(v or "").strip().upper()
    v = re.sub(r"^(NSE|BSE)\s*[:\-]\s*", "", v)
    v = re.sub(r"(-EQ|-BE|\.NS|\.BO|\s+EQ)$", "", v)
    return v.strip()


def parse_symbols(data: bytes, filename: str) -> list[str]:
    """Read share names from the first sheet of an .xlsx file or a .csv/.txt file.
    Uses the column whose header looks like Symbol/Stock/Share/Scrip/Name; otherwise the first column."""
    fn = (filename or "").lower()
    rows: list[list] = []
    if fn.endswith((".xlsx", ".xlsm")):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        for i, r in enumerate(ws.iter_rows(values_only=True)):
            rows.append(["" if c is None else str(c) for c in r])
            if i > 5000:
                break
    elif fn.endswith(".xls"):
        raise ValueError("Old .xls format: please save the file as .xlsx or .csv and import again")
    else:
        text = data.decode("utf-8-sig", errors="ignore")
        rows = [r for r in csv.reader(io.StringIO(text))][:5000]
    rows = [r for r in rows if any(str(c).strip() for c in r)]
    if not rows:
        return []
    col, start = 0, 0
    for j, h in enumerate(rows[0]):
        if HEADER_RX.search(str(h)):
            col, start = j, 1
            break
    out = []
    for r in rows[start:]:
        if col < len(r):
            v = _clean(r[col])
            if v and len(v) <= 25:
                out.append(v)
    return list(dict.fromkeys(out))
